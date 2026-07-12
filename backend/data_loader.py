"""
data_loader.py
Reads all 4 sheets from competitor_margin_analysis.xlsx and returns clean,
typed Python dicts/lists. All arithmetic was already computed in the spreadsheet;
we simply read the stored (computed) values — no LLM, no recalculation here.
"""

import os
from pathlib import Path
import openpyxl

EXCEL_PATH = Path(__file__).parent / "competitor_margin_analysis.xlsx"


def _load_workbook():
    return openpyxl.load_workbook(EXCEL_PATH, data_only=True)


# ---------------------------------------------------------------------------
# Sheet 1 – Price Delta  (columns A-J, rows 2-26, row 1 = header)
# A=SKU  B=Name  C=Category  D=Our Price  E=Competitor Price
# F=Delta%  G=Above10pct  H=Competitor  I=Source URL  J=Scraped At
# ---------------------------------------------------------------------------
def get_price_delta():
    wb = _load_workbook()
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=26, values_only=True):
        sku = row[0]
        if not sku:
            continue
        our_price = row[3]
        comp_price = row[4]
        delta_pct = row[5]
        above_flag = row[6]
        competitor = row[7]
        source_url = row[8]
        scraped_at = row[9]

        stale = comp_price is None

        rows.append({
            "sku":           str(sku).strip(),
            "name":          str(row[1]).strip() if row[1] else "",
            "category":      str(row[2]).strip() if row[2] else "",
            "our_price":     float(our_price) if our_price is not None else None,
            "comp_price":    float(comp_price) if comp_price is not None else None,
            "delta_pct":     round(float(delta_pct) * 100, 2) if delta_pct is not None else None,
            "above_10pct":   str(above_flag).strip() == "YES" if above_flag else False,
            "competitor":    str(competitor).strip() if competitor else "\u2014",
            "source_url":    str(source_url).strip() if source_url else None,
            "scraped_at":    str(scraped_at).strip() if scraped_at else None,
            "stale":         stale,
        })
    wb.close()
    return rows


# ---------------------------------------------------------------------------
# Sheet 2 – Margin Trend  (columns A-J, rows 2-26)
# A=SKU  B=Apr%  C=May%  D=Jun%  E-G=x1-3 (hidden)  H=Slope  I=Monotonic  J=Flag
# ---------------------------------------------------------------------------
def get_margin_trend():
    wb = _load_workbook()
    ws = wb.worksheets[1]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=26, values_only=True):
        sku = row[0]
        if not sku:
            continue
        apr = row[1]
        may = row[2]
        jun = row[3]
        slope = row[7]
        monotonic = row[8]
        flag = row[9]

        rows.append({
            "sku":       str(sku).strip(),
            "apr_pct":   round(float(apr) * 100, 2) if apr is not None else None,
            "may_pct":   round(float(may) * 100, 2) if may is not None else None,
            "jun_pct":   round(float(jun) * 100, 2) if jun is not None else None,
            "slope":     round(float(slope), 3) if slope is not None else None,
            "monotonic": str(monotonic).strip() if monotonic else "No",
            "declining": str(flag).strip() == "DECLINING" if flag else False,
        })
    wb.close()
    return rows


# ---------------------------------------------------------------------------
# Sheet 3 – Revenue at Risk  (columns A-K, rows 2-4)
# A=SKU  B=June Units  C=June Revenue  D=June Margin%  E=Current GP
# F=Implied Unit Cost  G=Comp Price  H=New Revenue  I=New GP  J=New Margin%  K=Profit Swing
# ---------------------------------------------------------------------------
def get_revenue_at_risk():
    wb = _load_workbook()
    ws = wb.worksheets[2]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=4, values_only=True):
        sku = row[0]
        if not sku:
            continue
        rows.append({
            "sku":             str(sku).strip(),
            "june_units":      int(row[1]) if row[1] is not None else None,
            "june_revenue":    round(float(row[2]), 2) if row[2] is not None else None,
            "june_margin_pct": round(float(row[3]) * 100, 2) if row[3] is not None else None,
            "current_gp":      round(float(row[4]), 2) if row[4] is not None else None,
            "implied_cost":    round(float(row[5]), 2) if row[5] is not None else None,
            "comp_price":      round(float(row[6]), 2) if row[6] is not None else None,
            "new_revenue":     round(float(row[7]), 2) if row[7] is not None else None,
            "new_gp":          round(float(row[8]), 2) if row[8] is not None else None,
            "new_margin_pct":  round(float(row[9]) * 100, 2) if row[9] is not None else None,
            "profit_swing":    round(float(row[10]), 2) if row[10] is not None else None,
        })
    wb.close()
    return rows


# ---------------------------------------------------------------------------
# Sheet 4 – Assumptions  (column A, rows 1-10)
# ---------------------------------------------------------------------------
def get_assumptions():
    wb = _load_workbook()
    ws = wb.worksheets[3]
    notes = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        cell = row[0]
        if cell and str(cell).strip():
            notes.append(str(cell).strip())
    wb.close()
    return notes


# ---------------------------------------------------------------------------
# Summary: aggregated stats for header cards
# ---------------------------------------------------------------------------
def get_summary(price_delta, margin_trend, revenue_at_risk):
    total_skus = len(price_delta)
    scraped_skus = sum(1 for r in price_delta if not r["stale"])
    declining_skus = sum(1 for r in margin_trend if r["declining"])
    above_market_skus = sum(1 for r in price_delta if r["above_10pct"])

    declining_ids = {r["sku"] for r in margin_trend if r["declining"]}
    above_ids = {r["sku"] for r in price_delta if r["above_10pct"]}
    join_skus = list(declining_ids & above_ids)

    total_profit_swing = sum(
        abs(r["profit_swing"]) for r in revenue_at_risk if r["profit_swing"] is not None
    )

    return {
        "total_skus":        total_skus,
        "scraped_skus":      scraped_skus,
        "declining_skus":    declining_skus,
        "above_market_skus": above_market_skus,
        "join_skus":         join_skus,
        "total_profit_swing": round(total_profit_swing, 2),
        "data_as_of":        "2026-07-12",
    }
